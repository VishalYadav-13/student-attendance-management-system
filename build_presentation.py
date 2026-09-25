import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor

# -----------------------------------------------------------------------------
# PRESENTATION CONFIGURATION & DESIGN SYSTEM (SAMS Vintage Academic Theme)
# Visual Identity: Calm + Distinguished + Professional + Academic (Sage & Ivory)
# -----------------------------------------------------------------------------
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]

# SAMS Academic Palette (aligned with academic-theme.css)
BG_DARK       = RGBColor(26, 31, 29)     # #1A1F1D SAMS Deep Calm Charcoal
CARD_BG       = RGBColor(36, 43, 40)     # #242B28 SAMS Forest Charcoal Surface
CARD_INNER    = RGBColor(30, 36, 34)     # #1E2422 SAMS Subtle Card Inner
CARD_BORDER   = RGBColor(56, 66, 62)     # #38423E SAMS Subtle Border
ACCENT_CYAN   = RGBColor(149, 178, 144)  # #95B290 SAMS Primary Accent: Sage Green
ACCENT_BLUE   = RGBColor(153, 179, 198)  # #99B3C6 SAMS Secondary Accent: Dusty Blue
ACCENT_GREEN  = RGBColor(111, 165, 116)  # #6FA574 SAMS Semantic Success Sage
ACCENT_INDIGO = RGBColor(126, 148, 168)  # #7E94A8 SAMS Academic Slate Blue
ACCENT_AMBER  = RGBColor(196, 133, 106)  # #C4856A SAMS Warm Accent: Terracotta
TEXT_LIGHT    = RGBColor(244, 239, 230)  # #F4EFE6 SAMS Warm Ivory / Soft Parchment
TEXT_MUTED    = RGBColor(203, 212, 207)  # #CBD4CF SAMS Warm Light Gray / Sage Tint
TEXT_DIM      = RGBColor(138, 153, 146)  # #8A9992 SAMS Dim Gray-Sage

FONT_NAME = "Segoe UI"

def create_base_slide():
    """Creates a slide with a clean dark background and standardized footer."""
    slide = prs.slides.add_slide(blank_layout)
    
    # Background fill
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_DARK
    bg.line.fill.background()
    
    # Top decorative accent line
    top_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0), Inches(11.733), Inches(0.06))
    top_line.fill.solid()
    top_line.fill.fore_color.rgb = ACCENT_CYAN
    top_line.line.fill.background()
    
    # Footer text
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(11.733), Inches(0.35))
    tf = footer_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Student Attendance Management System (SAMS)  |  Seminar & Project Initiation (CO5K)  |  Vishal Yadav"
    p.font.name = FONT_NAME
    p.font.size = Pt(9.5)
    p.font.color.rgb = TEXT_DIM
    
    return slide

def add_header(slide, section_tag, title_text, subtitle_text=None):
    """Adds a standardized academic header to the slide."""
    header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.35), Inches(11.733), Inches(1.0))
    tf = header_box.text_frame
    tf.word_wrap = True
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.margin_left = 0
    tf.margin_right = 0
    
    # Category tag
    p_tag = tf.paragraphs[0]
    p_tag.text = section_tag.upper()
    p_tag.font.name = FONT_NAME
    p_tag.font.size = Pt(9.5)
    p_tag.font.bold = True
    p_tag.font.color.rgb = ACCENT_CYAN
    p_tag.space_after = Pt(2)
    
    # Slide Title
    p_title = tf.add_paragraph()
    p_title.text = title_text
    p_title.font.name = FONT_NAME
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = TEXT_LIGHT
    
    if subtitle_text:
        p_sub = tf.add_paragraph()
        p_sub.text = subtitle_text
        p_sub.font.name = FONT_NAME
        p_sub.font.size = Pt(11)
        p_sub.font.color.rgb = TEXT_MUTED
        p_sub.space_before = Pt(2)

def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
    """Creates a card container shape with subtle border."""
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    card.fill.solid()
    card.fill.fore_color.rgb = bg_color
    card.line.color.rgb = border_color
    card.line.width = Pt(1)
    return card

# -----------------------------------------------------------------------------
# SLIDE 1: PROJECT TITLE
# -----------------------------------------------------------------------------
def build_slide_1():
    slide = prs.slides.add_slide(blank_layout)
    
    # Full background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG_DARK
    bg.line.fill.background()
    
    # Right visual card container
    glow_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.2), Inches(1.1), Inches(4.333), Inches(5.4))
    glow_card.fill.solid()
    glow_card.fill.fore_color.rgb = CARD_BG
    glow_card.line.color.rgb = CARD_BORDER
    glow_card.line.width = Pt(1.5)
    
    # Inner Tech Badge inside right card
    ai_badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.55), Inches(1.35), Inches(3.633), Inches(0.55))
    ai_badge.fill.solid()
    ai_badge.fill.fore_color.rgb = CARD_INNER
    ai_badge.line.color.rgb = ACCENT_CYAN
    tf_badge = ai_badge.text_frame
    p = tf_badge.paragraphs[0]
    p.text = "✦  EMERGING TECH: AI & COMPUTER VISION"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_NAME
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    
    pillars = [
        ("AI Face Recognition Engine", "CompreFace biometric landmark detection & server-side verification"),
        ("Centralized Cloud Architecture", "Vercel edge static frontend + Render containerized PHP 8.2 backend"),
        ("Relational Data Management", "PostgreSQL database with strict RBAC, audit trails & unique constraints"),
        ("Automated Academic Workflows", "Teacher session studio, instant roll-call, and compliance analytics")
    ]
    
    y_offset = Inches(2.05)
    for title, desc in pillars:
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.55), y_offset, Inches(3.633), Inches(0.95))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_INNER
        box.line.color.rgb = CARD_BORDER
        tf_b = box.text_frame
        tf_b.word_wrap = True
        tf_b.margin_top = Pt(6)
        tf_b.margin_left = Pt(10)
        tf_b.margin_right = Pt(10)
        
        p1 = tf_b.paragraphs[0]
        p1.text = "• " + title
        p1.font.name = FONT_NAME
        p1.font.size = Pt(11)
        p1.font.bold = True
        p1.font.color.rgb = ACCENT_CYAN
        
        p2 = tf_b.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_NAME
        p2.font.size = Pt(9.5)
        p2.font.color.rgb = TEXT_MUTED
        
        y_offset += Inches(1.08)
    
    # Left Content Area (Title & Presentation Details)
    inst_badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.1), Inches(6.8), Inches(0.42))
    inst_badge.fill.solid()
    inst_badge.fill.fore_color.rgb = CARD_INNER
    inst_badge.line.color.rgb = ACCENT_CYAN
    tf_inst = inst_badge.text_frame
    p_inst = tf_inst.paragraphs[0]
    p_inst.text = "MSBTE DIPLOMA IN COMPUTER ENGINEERING  |  SEMINAR & PROJECT INITIATION"
    p_inst.font.name = FONT_NAME
    p_inst.font.size = Pt(9.5)
    p_inst.font.bold = True
    p_inst.font.color.rgb = ACCENT_CYAN
    p_inst.alignment = PP_ALIGN.CENTER
    
    # Main Project Title Box
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(7.0), Inches(2.3))
    tf_title = title_box.text_frame
    tf_title.word_wrap = True
    tf_title.margin_left = 0
    tf_title.margin_top = 0
    
    p_maintitle = tf_title.paragraphs[0]
    p_maintitle.text = "STUDENT ATTENDANCE\nMANAGEMENT SYSTEM (SAMS)"
    p_maintitle.font.name = FONT_NAME
    p_maintitle.font.size = Pt(28)
    p_maintitle.font.bold = True
    p_maintitle.font.color.rgb = TEXT_LIGHT
    p_maintitle.space_after = Pt(10)
    
    p_sub = tf_title.add_paragraph()
    p_sub.text = "Subject: Seminar & Project Initiation\nTopic: Emerging Trends in Technology"
    p_sub.font.name = FONT_NAME
    p_sub.font.size = Pt(14)
    p_sub.font.bold = True
    p_sub.font.color.rgb = ACCENT_BLUE
    
    # Student Presentation Metadata Card
    meta_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(4.35), Inches(7.0), Inches(2.15))
    meta_card.fill.solid()
    meta_card.fill.fore_color.rgb = CARD_BG
    meta_card.line.color.rgb = CARD_BORDER
    
    # Two-column layout inside meta_card for labels and values
    lbl_box = slide.shapes.add_textbox(Inches(1.1), Inches(4.5), Inches(2.2), Inches(1.8))
    tf_lbl = lbl_box.text_frame
    tf_lbl.word_wrap = True
    tf_lbl.margin_left = 0
    tf_lbl.margin_top = 0
    
    val_box = slide.shapes.add_textbox(Inches(3.4), Inches(4.5), Inches(4.2), Inches(1.8))
    tf_val = val_box.text_frame
    tf_val.word_wrap = True
    tf_val.margin_left = 0
    tf_val.margin_top = 0
    
    meta_items = [
        ("Presented by:", "Vishal Yadav"),
        ("Class / Branch:", "CO5K  (Computer Engineering)"),
        ("Enrollment No.:", "________________"),
        ("Academic Year:", "2026 – 2027")
    ]
    
    for i, (label, val) in enumerate(meta_items):
        p_l = tf_lbl.paragraphs[0] if i == 0 else tf_lbl.add_paragraph()
        p_l.text = label
        p_l.font.name = FONT_NAME
        p_l.font.size = Pt(12)
        p_l.font.bold = True
        p_l.font.color.rgb = ACCENT_CYAN
        p_l.space_after = Pt(5)
        
        p_v = tf_val.paragraphs[0] if i == 0 else tf_val.add_paragraph()
        p_v.text = val
        p_v.font.name = FONT_NAME
        p_v.font.size = Pt(12)
        p_v.font.bold = (label == "Presented by:" or label == "Enrollment No.:")
        p_v.font.color.rgb = TEXT_LIGHT
        p_v.space_after = Pt(5)

# -----------------------------------------------------------------------------
# SLIDE 2: 1. INTRODUCTION
# -----------------------------------------------------------------------------
def build_slide_2():
    slide = create_base_slide()
    add_header(slide, "1. Introduction", "1. Introduction", "Overview of Academic Attendance Management & SAMS Vision")
    
    # Left Column: Key Points Card
    card_left = add_card(slide, Inches(0.8), Inches(1.6), Inches(7.5), Inches(5.1))
    tf = card_left.text_frame
    tf.word_wrap = True
    tf.margin_left = Pt(18)
    tf.margin_top = Pt(18)
    tf.margin_right = Pt(18)
    
    points = [
        ("Indispensable Academic Process", "Attendance is an important part of academic management, directly influencing student discipline, accreditation, and eligibility."),
        ("Manual Effort Challenges", "Traditional attendance methods require continuous manual effort from teachers for roll-call and record-keeping."),
        ("Instructional Time Loss", "Manual roll-call consumes valuable classroom lecture time and remains susceptible to human recording errors."),
        ("Centralized Digital Management", "SAMS provides a centralized digital system for efficiently managing student attendance across institutes."),
        ("Comprehensive Academic Workflows", "The system supports complete student management, teacher workflows, attendance recording, and analytics."),
        ("AI Face Recognition Integration", "The project explores the use of AI-based face recognition to automate student attendance verification."),
        ("Web-Based Modern Accessibility", "Designed as a modern web-based application, accessible across campus devices without specialized client installs.")
    ]
    
    for i, (headline, body) in enumerate(points):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"• {headline}: "
        p.font.name = FONT_NAME
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = ACCENT_CYAN
        
        run = p.add_run()
        run.text = body
        run.font.name = FONT_NAME
        run.font.size = Pt(10)
        run.font.bold = False
        run.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(6)
        
    # Right Column: Visual Diagram Card
    card_right = add_card(slide, Inches(8.6), Inches(1.6), Inches(3.933), Inches(5.1))
    
    # Title badge for right column
    title_badge = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.85), Inches(1.85), Inches(3.433), Inches(0.45))
    title_badge.fill.solid()
    title_badge.fill.fore_color.rgb = CARD_INNER
    title_badge.line.color.rgb = ACCENT_CYAN
    tf_tb = title_badge.text_frame
    p_tb = tf_tb.paragraphs[0]
    p_tb.text = "SAMS OPERATIONAL FLOW"
    p_tb.alignment = PP_ALIGN.CENTER
    p_tb.font.name = FONT_NAME
    p_tb.font.size = Pt(10.5)
    p_tb.font.bold = True
    p_tb.font.color.rgb = ACCENT_CYAN
    
    flow_steps = [
        ("1. Faculty / Teacher", "Creates lecture session & selects class division", ACCENT_INDIGO),
        ("2. SAMS Platform", "Loads active roster & coordinates verification", ACCENT_CYAN),
        ("3. Student Attendance", "Recorded via AI Face Recognition or Manual Grid", ACCENT_GREEN),
        ("4. Reports & Analytics", "Stored in PostgreSQL & available for export", ACCENT_BLUE)
    ]
    
    y_flow = Inches(2.45)
    for title, desc, color in flow_steps:
        step_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.85), y_flow, Inches(3.433), Inches(0.78))
        step_box.fill.solid()
        step_box.fill.fore_color.rgb = CARD_INNER
        step_box.line.color.rgb = color
        step_box.line.width = Pt(1.5)
        
        tf_s = step_box.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = Pt(10)
        tf_s.margin_top = Pt(6)
        p1 = tf_s.paragraphs[0]
        p1.text = title
        p1.font.name = FONT_NAME
        p1.font.size = Pt(10.5)
        p1.font.bold = True
        p1.font.color.rgb = color
        
        p2 = tf_s.add_paragraph()
        p2.text = desc
        p2.font.name = FONT_NAME
        p2.font.size = Pt(9)
        p2.font.color.rgb = TEXT_MUTED
        
        y_flow += Inches(0.86)
        
        # Add arrow if not last
        if title != flow_steps[-1][0]:
            arrow = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(10.42), y_flow - Inches(0.08), Inches(0.3), Inches(0.18))
            arrow.fill.solid()
            arrow.fill.fore_color.rgb = ACCENT_CYAN
            arrow.line.fill.background()
            y_flow += Inches(0.18)

# -----------------------------------------------------------------------------
# SLIDE 3: 2. EMERGING TECHNOLOGY USED
# -----------------------------------------------------------------------------
def build_slide_3():
    slide = create_base_slide()
    add_header(slide, "2. Emerging Technology Used", "2. Emerging Technology Used", "Core Technological Pillars & Rationale for Adoption")
    
    tech_pillars = [
        ("Artificial Intelligence / Computer Vision", 
         ["Face detection and landmark extraction", 
          "Biometric identity verification", 
          "Automated attendance verification", 
          "Reduces human verification delays"], 
         ACCENT_CYAN, Inches(0.8), Inches(1.55), Inches(5.7), Inches(2.2)),
         
        ("Cloud Computing Architecture", 
         ["Cloud-hosted frontend (Vercel CDN edge)", 
          "Cloud-hosted backend (Render container)", 
          "Remote managed PostgreSQL database", 
          "Accessible across entire campus network"], 
         ACCENT_INDIGO, Inches(6.8), Inches(1.55), Inches(5.733), Inches(2.2)),
         
        ("Modern Web Technologies", 
         ["Frontend: HTML5, CSS3, Vanilla JS (ES6+)", 
          "Backend: PHP 8.2+ RESTful routing", 
          "Database: PostgreSQL 13+ relational engine", 
          "Stateless JWT tokens & secure session guards"], 
         ACCENT_BLUE, Inches(0.8), Inches(3.9), Inches(5.7), Inches(2.1)),
         
        ("Face Recognition Service (CompreFace)", 
         ["CompreFace open-source recognition engine", 
          "Server-side face recognition architecture", 
          "Secure mathematical vector matching", 
          "High accuracy without client credential exposure"], 
         ACCENT_GREEN, Inches(6.8), Inches(3.9), Inches(5.733), Inches(2.1))
    ]
    
    for title, items, color, left, top, width, height in tech_pillars:
        card = add_card(slide, left, top, width, height)
        card.line.color.rgb = color
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Pt(14)
        tf.margin_top = Pt(10)
        
        p_title = tf.paragraphs[0]
        p_title.text = title
        p_title.font.name = FONT_NAME
        p_title.font.size = Pt(12)
        p_title.font.bold = True
        p_title.font.color.rgb = color
        p_title.space_after = Pt(4)
        
        for item in items:
            p = tf.add_paragraph()
            p.text = "• " + item
            p.font.name = FONT_NAME
            p.font.size = Pt(10)
            p.font.color.rgb = TEXT_LIGHT
            p.space_after = Pt(2)
            
    # Bottom Bar: Why these technologies are used
    why_bar = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.15), Inches(11.733), Inches(0.75))
    why_bar.fill.solid()
    why_bar.fill.fore_color.rgb = CARD_INNER
    why_bar.line.color.rgb = ACCENT_CYAN
    tf_why = why_bar.text_frame
    tf_why.word_wrap = True
    tf_why.margin_left = Pt(14)
    tf_why.margin_top = Pt(8)
    
    p_whytitle = tf_why.paragraphs[0]
    p_whytitle.text = "WHY THESE TECHNOLOGIES?  "
    p_whytitle.font.name = FONT_NAME
    p_whytitle.font.size = Pt(10.5)
    p_whytitle.font.bold = True
    p_whytitle.font.color.rgb = ACCENT_CYAN
    
    run_why = p_whytitle.add_run()
    run_why.text = "Reduces manual attendance work  |  Improves automation  |  Faster identity verification  |  Centralized digital records  |  Universal web accessibility"
    run_why.font.name = FONT_NAME
    run_why.font.size = Pt(10)
    run_why.font.bold = False
    run_why.font.color.rgb = TEXT_LIGHT

# -----------------------------------------------------------------------------
# SLIDE 4: 3. PROBLEM STATEMENT
# -----------------------------------------------------------------------------
def build_slide_4():
    slide = create_base_slide()
    add_header(slide, "3. Problem Statement", "3. Problem Statement", "Comparison: Existing Manual System vs. Proposed SAMS Solution")
    
    # Left Card: Existing System
    left_card = add_card(slide, Inches(0.8), Inches(1.6), Inches(5.6), Inches(4.3))
    left_card.line.color.rgb = ACCENT_AMBER
    tf_l = left_card.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = Pt(16)
    tf_l.margin_top = Pt(14)
    
    p_lh = tf_l.paragraphs[0]
    p_lh.text = "EXISTING SYSTEM  (Traditional / Manual)"
    p_lh.font.name = FONT_NAME
    p_lh.font.size = Pt(13)
    p_lh.font.bold = True
    p_lh.font.color.rgb = ACCENT_AMBER
    p_lh.space_after = Pt(8)
    
    existing_points = [
        "Attendance is recorded manually on paper registers or loose sheets.",
        "Teachers spend 10–15 minutes of precious classroom time taking attendance.",
        "Manual roll-call and manual entry frequently lead to human recording errors.",
        "Searching historical attendance records and compiling monthly totals is tedious.",
        "Generating compliance and semester reports requires extensive manual effort.",
        "Manual attendance does not provide automated identity verification or proxy prevention."
    ]
    for pt in existing_points:
        p = tf_l.add_paragraph()
        p.text = "✖  " + pt
        p.font.name = FONT_NAME
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(4)
        
    # Right Card: Proposed System
    right_card = add_card(slide, Inches(6.8), Inches(1.6), Inches(5.733), Inches(4.3))
    right_card.line.color.rgb = ACCENT_GREEN
    tf_r = right_card.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = Pt(16)
    tf_r.margin_top = Pt(14)
    
    p_rh = tf_r.paragraphs[0]
    p_rh.text = "PROPOSED SYSTEM  (SAMS Centralized Platform)"
    p_rh.font.name = FONT_NAME
    p_rh.font.size = Pt(13)
    p_rh.font.bold = True
    p_rh.font.color.rgb = ACCENT_GREEN
    p_rh.space_after = Pt(8)
    
    proposed_points = [
        "Centralized web-based Student Attendance Management System.",
        "Structured PostgreSQL database eliminating paper registers completely.",
        "Streamlined teacher-based attendance workflow saving lecture hours.",
        "Automated AI face verification option for instant, contactless identity check.",
        "Instant attendance percentage reports, analytics, and 75% shortage alarms.",
        "Role-based access control (RBAC) securely separating Admin and Teacher privileges."
    ]
    for pt in proposed_points:
        p = tf_r.add_paragraph()
        p.text = "✔  " + pt
        p.font.name = FONT_NAME
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(4)
        
    # Bottom Comparison Banner: Existing vs Proposed
    banner = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(6.05), Inches(11.733), Inches(0.85))
    banner.fill.solid()
    banner.fill.fore_color.rgb = CARD_INNER
    banner.line.color.rgb = CARD_BORDER
    tf_b = banner.text_frame
    tf_b.word_wrap = True
    tf_b.margin_top = Pt(10)
    
    p_b = tf_b.paragraphs[0]
    p_b.alignment = PP_ALIGN.CENTER
    
    run1 = p_b.add_run()
    run1.text = "EXISTING: Manual  ➔  Slow  ➔  Error-Prone  ➔  Paper Registers"
    run1.font.name = FONT_NAME
    run1.font.size = Pt(11)
    run1.font.bold = True
    run1.font.color.rgb = ACCENT_AMBER
    
    run_vs = p_b.add_run()
    run_vs.text = "     VS     "
    run_vs.font.name = FONT_NAME
    run_vs.font.size = Pt(11)
    run_vs.font.bold = True
    run_vs.font.color.rgb = TEXT_DIM
    
    run2 = p_b.add_run()
    run2.text = "PROPOSED: Digital  ➔  Automated  ➔  Centralized  ➔  Verified Records"
    run2.font.name = FONT_NAME
    run2.font.size = Pt(11)
    run2.font.bold = True
    run2.font.color.rgb = ACCENT_GREEN

# -----------------------------------------------------------------------------
# SLIDE 5: 4. OBJECTIVES
# -----------------------------------------------------------------------------
def build_slide_5():
    slide = create_base_slide()
    add_header(slide, "4. Objectives", "4. Objectives", "Core Technical & Academic Goals of the Project")
    
    objectives = [
        ("1", "Centralized System", "Develop a centralized web-based student attendance management system for academic institutes."),
        ("2", "Reduce Manual Effort", "Significantly reduce manual roll-call effort and eliminate paper-based record maintenance."),
        ("3", "Role-Based Security", "Provide secure student and teacher management with role-based access control (RBAC)."),
        ("4", "Structured Database", "Maintain institutional attendance records in a secure, structured PostgreSQL relational database."),
        ("5", "Reliable Manual Workflow", "Provide an intuitive manual grid attendance marking workflow as a reliable standard and backup."),
        ("6", "AI Face Verification", "Integrate AI-based face recognition to enable automated, contactless student attendance verification."),
        ("7", "Reports & Statistics", "Generate comprehensive attendance reports, subject summaries, and student shortage statistics."),
        ("8", "Accuracy & Efficiency", "Improve institutional efficiency and eliminate human errors and proxy attendance risks."),
        ("9", "Future Foundation", "Establish a scalable technical foundation for future intelligent academic campus features.")
    ]
    
    cols = 3
    col_w = Inches(3.75)
    row_h = Inches(1.55)
    start_x = Inches(0.8)
    start_y = Inches(1.6)
    gap_x = Inches(0.24)
    gap_y = Inches(0.2)
    
    for idx, (num, title, desc) in enumerate(objectives):
        r = idx // cols
        c = idx % cols
        x = start_x + c * (col_w + gap_x)
        y = start_y + r * (row_h + gap_y)
        
        card = add_card(slide, x, y, col_w, row_h)
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Pt(12)
        tf.margin_top = Pt(10)
        tf.margin_right = Pt(12)
        
        p = tf.paragraphs[0]
        p.text = f"[{num}] {title}"
        p.font.name = FONT_NAME
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = ACCENT_CYAN
        p.space_after = Pt(3)
        
        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = FONT_NAME
        p_desc.font.size = Pt(9.5)
        p_desc.font.color.rgb = TEXT_LIGHT

# -----------------------------------------------------------------------------
# SLIDE 6: 5. PROPOSED SYSTEM
# -----------------------------------------------------------------------------
def build_slide_6():
    slide = create_base_slide()
    add_header(slide, "5. Proposed System", "5. Proposed System", "High-Level Architecture & End-to-End System Organization")
    
    pipeline_steps = [
        ("Admin / Teacher", "Authorized User", ACCENT_INDIGO),
        ("Web Interface", "HTML / CSS / JS", ACCENT_BLUE),
        ("SAMS Backend", "PHP 8.2+ API", ACCENT_CYAN),
        ("Auth & Security", "JWT & RBAC", ACCENT_CYAN),
        ("Attendance Logic", "Session Engine", ACCENT_GREEN),
        ("AI Face Engine", "CompreFace AI", ACCENT_GREEN),
        ("PostgreSQL DB", "Relational Data", ACCENT_BLUE),
        ("Reports & Analytics", "Visual Results", ACCENT_INDIGO)
    ]
    
    pipe_w = Inches(1.36)
    pipe_h = Inches(1.25)
    gap = Inches(0.12)
    start_x = Inches(0.8)
    y_pipe = Inches(1.55)
    
    for i, (title, subtitle, color) in enumerate(pipeline_steps):
        x = start_x + i * (pipe_w + gap)
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y_pipe, pipe_w, pipe_h)
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_INNER
        box.line.color.rgb = color
        box.line.width = Pt(1.5)
        
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = Pt(4)
        tf.margin_right = Pt(4)
        tf.margin_top = Pt(8)
        
        p_num = tf.paragraphs[0]
        p_num.text = f"Step {i+1}"
        p_num.alignment = PP_ALIGN.CENTER
        p_num.font.name = FONT_NAME
        p_num.font.size = Pt(8.5)
        p_num.font.bold = True
        p_num.font.color.rgb = color
        
        p_title = tf.add_paragraph()
        p_title.text = title
        p_title.alignment = PP_ALIGN.CENTER
        p_title.font.name = FONT_NAME
        p_title.font.size = Pt(10)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_LIGHT
        
        p_sub = tf.add_paragraph()
        p_sub.text = subtitle
        p_sub.alignment = PP_ALIGN.CENTER
        p_sub.font.name = FONT_NAME
        p_sub.font.size = Pt(8)
        p_sub.font.color.rgb = TEXT_MUTED

    cards_data = [
        ("Administrator Control & Governance", 
         ["Centralized management of departments, courses, teachers, and students.",
          "Configures academic calendars, class divisions, and attendance thresholds.",
          "Maintains immutable audit logs of system activity and teacher overrides."],
         ACCENT_INDIGO, Inches(0.8), Inches(3.1), Inches(3.75), Inches(3.7)),
         
        ("Teacher Studio & Attendance Workflows", 
         ["Creates real-time lecture and laboratory attendance sessions.",
          "Dynamically retrieves class and division student rosters.",
          "Provides manual attendance grid with instant batch marking options.",
          "Performs continuous or individual AI face verification seamlessly."],
         ACCENT_CYAN, Inches(4.79), Inches(3.1), Inches(3.75), Inches(3.7)),
         
        ("Data Persistence & Reporting Engine", 
         ["Enrolled student biometric vectors mapped securely to academic records.",
          "PostgreSQL relational storage prevents duplicate entries per session.",
          "Calculates cumulative attendance and triggers 75% shortage warnings.",
          "Generates exportable compliance reports for HODs and administration."],
         ACCENT_GREEN, Inches(8.78), Inches(3.1), Inches(3.75), Inches(3.7))
    ]
    
    for title, points, color, left, top, width, height in cards_data:
        c = add_card(slide, left, top, width, height)
        c.line.color.rgb = color
        tf = c.text_frame
        tf.word_wrap = True
        tf.margin_left = Pt(14)
        tf.margin_top = Pt(14)
        tf.margin_right = Pt(14)
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_NAME
        p.font.size = Pt(11.5)
        p.font.bold = True
        p.font.color.rgb = color
        p.space_after = Pt(8)
        
        for pt in points:
            p_pt = tf.add_paragraph()
            p_pt.text = "• " + pt
            p_pt.font.name = FONT_NAME
            p_pt.font.size = Pt(10)
            p_pt.font.color.rgb = TEXT_LIGHT
            p_pt.space_after = Pt(6)

# -----------------------------------------------------------------------------
# SLIDE 7: 6. WORKING PRINCIPLE
# -----------------------------------------------------------------------------
def build_slide_7():
    slide = create_base_slide()
    add_header(slide, "6. Working Principle", "6. Working Principle", "Step-by-Step Attendance Execution Cycle & Workflow")
    
    flow_steps = [
        "1. Login", "2. Session", "3. Input", "4. Processing", 
        "5. Verification", "6. Database", "7. Result"
    ]
    sub_labels = [
        "Auth User", "Select Class", "Camera / Grid", "CompreFace AI", 
        "Match Identity", "PostgreSQL", "Report / KPI"
    ]
    
    w_box = Inches(1.58)
    h_box = Inches(0.85)
    gap = Inches(0.11)
    y_flow = Inches(1.55)
    
    for i in range(7):
        x = Inches(0.8) + i * (w_box + gap)
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y_flow, w_box, h_box)
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_INNER
        box.line.color.rgb = ACCENT_CYAN
        
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_top = Pt(6)
        
        p = tf.paragraphs[0]
        p.text = flow_steps[i]
        p.alignment = PP_ALIGN.CENTER
        p.font.name = FONT_NAME
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE
        
        p2 = tf.add_paragraph()
        p2.text = sub_labels[i]
        p2.alignment = PP_ALIGN.CENTER
        p2.font.name = FONT_NAME
        p2.font.size = Pt(8.5)
        p2.font.color.rgb = TEXT_MUTED

    card_l = add_card(slide, Inches(0.8), Inches(2.65), Inches(5.7), Inches(4.2))
    tf_l = card_l.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = Pt(14)
    tf_l.margin_top = Pt(12)
    tf_l.margin_right = Pt(14)
    
    p_lh = tf_l.paragraphs[0]
    p_lh.text = "PHASE 1: AUTHENTICATION & SESSION SETUP"
    p_lh.font.name = FONT_NAME
    p_lh.font.size = Pt(11)
    p_lh.font.bold = True
    p_lh.font.color.rgb = ACCENT_CYAN
    p_lh.space_after = Pt(6)
    
    steps_1_to_6 = [
        ("1. User Login", "Teacher or Administrator logs into the SAMS web application."),
        ("2. Role Verification", "Backend verifies user credentials, role permissions, and session tokens."),
        ("3. Session Selection", "Teacher selects academic course, semester, division, and subject."),
        ("4. Roster Retrieval", "System dynamically loads the active enrolled student roster from PostgreSQL."),
        ("5. Manual Attendance", "Teacher can record attendance manually using the interactive grid studio."),
        ("6. Camera Capture", "For face attendance, browser camera captures the student's face via Web API.")
    ]
    for num_title, desc in steps_1_to_6:
        p = tf_l.add_paragraph()
        p.text = f"{num_title}: "
        p.font.name = FONT_NAME
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE
        run = p.add_run()
        run.text = desc
        run.font.name = FONT_NAME
        run.font.size = Pt(9.5)
        run.font.bold = False
        run.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(4)

    card_r = add_card(slide, Inches(6.8), Inches(2.65), Inches(5.733), Inches(4.2))
    tf_r = card_r.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = Pt(14)
    tf_r.margin_top = Pt(12)
    tf_r.margin_right = Pt(14)
    
    p_rh = tf_r.paragraphs[0]
    p_rh.text = "PHASE 2: AI RECOGNITION & RECORD PERSISTENCE"
    p_rh.font.name = FONT_NAME
    p_rh.font.size = Pt(11)
    p_rh.font.bold = True
    p_rh.font.color.rgb = ACCENT_GREEN
    p_rh.space_after = Pt(6)
    
    steps_7_to_11 = [
        ("7. Face Processing", "Captured frame is analyzed by CompreFace face recognition service."),
        ("8. Identity Verification", "Extracted facial landmarks are mapped to enrolled student profiles."),
        ("9. Backend Validation", "PHP backend validates session ID, timetable, and duplicate marking rules."),
        ("10. Database Persistence", "Confirmed attendance record is permanently saved in PostgreSQL."),
        ("11. Report Generation", "Attendance percentages update in real-time for student & faculty dashboards.")
    ]
    for num_title, desc in steps_7_to_11:
        p = tf_r.add_paragraph()
        p.text = f"{num_title}: "
        p.font.name = FONT_NAME
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = ACCENT_GREEN
        run = p.add_run()
        run.text = desc
        run.font.name = FONT_NAME
        run.font.size = Pt(9.5)
        run.font.bold = False
        run.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(6)

# -----------------------------------------------------------------------------
# SLIDE 8: 7. H/W AND S/W REQUIREMENTS
# -----------------------------------------------------------------------------
def build_slide_8():
    slide = create_base_slide()
    add_header(slide, "7. H/W and S/W Requirements", "7. H/W and S/W Requirements", "Realistic Specifications for Full-Stack Web Development & Deployment")
    
    card_hw = add_card(slide, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.1))
    card_hw.line.color.rgb = ACCENT_CYAN
    tf_hw = card_hw.text_frame
    tf_hw.word_wrap = True
    tf_hw.margin_left = Pt(16)
    tf_hw.margin_top = Pt(14)
    tf_hw.margin_right = Pt(16)
    
    p_hwh = tf_hw.paragraphs[0]
    p_hwh.text = "HARDWARE REQUIREMENTS (Web-Based System)"
    p_hwh.font.name = FONT_NAME
    p_hwh.font.size = Pt(12.5)
    p_hwh.font.bold = True
    p_hwh.font.color.rgb = ACCENT_CYAN
    p_hwh.space_after = Pt(8)
    
    hw_specs = [
        ("Client Device", "Standard PC, Laptop, or Tablet for faculty/admin operation."),
        ("Camera", "Standard HD Webcam (720p / 1080p) or integrated laptop camera."),
        ("Processor", "Dual-core CPU (Intel Core i3 / AMD Ryzen 3 or higher)."),
        ("System Memory (RAM)", "Minimum 4 GB RAM recommended for smooth development & runtime."),
        ("Storage", "Standard 20 GB free disk space for codebase, Docker, & database."),
        ("Network Connectivity", "Active LAN / Wi-Fi / Broadband Internet connection."),
        ("Hardware Distinction", "Pure software web project: No Arduino, microcontrollers, or IoT sensors needed.")
    ]
    for comp, req in hw_specs:
        p = tf_hw.add_paragraph()
        p.text = f"• {comp}: "
        p.font.name = FONT_NAME
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT
        run = p.add_run()
        run.text = req
        run.font.name = FONT_NAME
        run.font.size = Pt(10)
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(6)

    card_sw = add_card(slide, Inches(6.8), Inches(1.6), Inches(5.733), Inches(5.1))
    card_sw.line.color.rgb = ACCENT_GREEN
    tf_sw = card_sw.text_frame
    tf_sw.word_wrap = True
    tf_sw.margin_left = Pt(16)
    tf_sw.margin_top = Pt(14)
    tf_sw.margin_right = Pt(16)
    
    p_swh = tf_sw.paragraphs[0]
    p_swh.text = "SOFTWARE REQUIREMENTS"
    p_swh.font.name = FONT_NAME
    p_swh.font.size = Pt(12.5)
    p_swh.font.bold = True
    p_swh.font.color.rgb = ACCENT_GREEN
    p_swh.space_after = Pt(8)
    
    sw_specs = [
        ("Frontend Technologies", "HTML5, CSS3, Modern Vanilla JavaScript (ES6+), Chart.js."),
        ("Backend Framework", "PHP 8.2+ with REST router, JSON responses, and PDO drivers."),
        ("Database Management", "PostgreSQL 13+ relational database with foreign key cascades."),
        ("Face Recognition", "CompreFace Facial Recognition Engine (Docker-based microservice)."),
        ("Containerization", "Docker & Docker Compose for service orchestration."),
        ("Cloud Deployment", "Vercel (Frontend Edge Hosting) & Render (PHP Backend & DB)."),
        ("Development Tools", "VS Code / Antigravity IDE, Git & GitHub for Version Control."),
        ("Client Browser", "Modern Web Browser (Google Chrome, Microsoft Edge, Firefox).")
    ]
    for cat, tools in sw_specs:
        p = tf_sw.add_paragraph()
        p.text = f"• {cat}: "
        p.font.name = FONT_NAME
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = TEXT_LIGHT
        run = p.add_run()
        run.text = tools
        run.font.name = FONT_NAME
        run.font.size = Pt(10)
        run.font.bold = False
        run.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(4)

# -----------------------------------------------------------------------------
# SLIDE 9: 8. BLOCK DIAGRAM
# -----------------------------------------------------------------------------
def build_slide_9():
    slide = create_base_slide()
    add_header(slide, "8. Block Diagram", "8. Block Diagram", "Architectural Block Diagram & Cloud Deployment Infrastructure")
    
    # Block Diagram Container Card (Top)
    top_container = add_card(slide, Inches(0.8), Inches(1.48), Inches(11.733), Inches(3.82))
    
    # 1. User Layer Box
    b_user = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(4.5), Inches(1.6), Inches(4.333), Inches(0.44))
    b_user.fill.solid()
    b_user.fill.fore_color.rgb = CARD_INNER
    b_user.line.color.rgb = ACCENT_INDIGO
    b_user.line.width = Pt(1.5)
    p = b_user.text_frame.paragraphs[0]
    p.text = "ADMINISTRATOR  /  TEACHER"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_NAME
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = ACCENT_INDIGO
    
    # Down arrow 1
    arr1 = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.52), Inches(2.07), Inches(0.28), Inches(0.14))
    arr1.fill.solid()
    arr1.fill.fore_color.rgb = ACCENT_CYAN
    arr1.line.fill.background()
    
    # 2. Web Application Box
    b_web = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.8), Inches(2.23), Inches(5.733), Inches(0.44))
    b_web.fill.solid()
    b_web.fill.fore_color.rgb = CARD_INNER
    b_web.line.color.rgb = ACCENT_CYAN
    b_web.line.width = Pt(1.5)
    p = b_web.text_frame.paragraphs[0]
    p.text = "WEB APPLICATION  (HTML5 / CSS3 / JavaScript / Camera API)"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_NAME
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    
    # Down arrow 2
    arr2 = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.52), Inches(2.70), Inches(0.28), Inches(0.14))
    arr2.fill.solid()
    arr2.fill.fore_color.rgb = ACCENT_CYAN
    arr2.line.fill.background()
    
    # 3. PHP Backend Box
    b_php = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.5), Inches(2.86), Inches(6.333), Inches(0.46))
    b_php.fill.solid()
    b_php.fill.fore_color.rgb = CARD_INNER
    b_php.line.color.rgb = ACCENT_BLUE
    b_php.line.width = Pt(1.5)
    p = b_php.text_frame.paragraphs[0]
    p.text = "PHP BACKEND  (Authentication & Attendance Business Logic)"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_NAME
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE
    
    # Left Branch Arrow (to PostgreSQL)
    arr_l = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(3.6), Inches(3.34), Inches(0.28), Inches(0.14))
    arr_l.fill.solid()
    arr_l.fill.fore_color.rgb = ACCENT_GREEN
    arr_l.line.fill.background()
    
    # Right Branch Arrow (to CompreFace)
    arr_r = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(9.4), Inches(3.34), Inches(0.28), Inches(0.14))
    arr_r.fill.solid()
    arr_r.fill.fore_color.rgb = ACCENT_CYAN
    arr_r.line.fill.background()
    
    # Left Branch: PostgreSQL DB
    b_db = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(3.50), Inches(4.5), Inches(0.68))
    b_db.fill.solid()
    b_db.fill.fore_color.rgb = CARD_INNER
    b_db.line.color.rgb = ACCENT_GREEN
    b_db.line.width = Pt(1.5)
    tf_db = b_db.text_frame
    p = tf_db.paragraphs[0]
    p.text = "PostgreSQL Relational Database"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_NAME
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p2 = tf_db.add_paragraph()
    p2.text = "Users, Classes, Attendance Sessions & Records"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.name = FONT_NAME
    p2.font.size = Pt(8.5)
    p2.font.color.rgb = TEXT_MUTED
    
    # Right Branch: AI Face Recognition (CompreFace)
    b_ai = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.333), Inches(3.50), Inches(4.5), Inches(0.68))
    b_ai.fill.solid()
    b_ai.fill.fore_color.rgb = CARD_INNER
    b_ai.line.color.rgb = ACCENT_CYAN
    b_ai.line.width = Pt(1.5)
    tf_ai = b_ai.text_frame
    p = tf_ai.paragraphs[0]
    p.text = "AI Face Recognition (CompreFace)"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_NAME
    p.font.size = Pt(10.5)
    p.font.bold = True
    p.font.color.rgb = ACCENT_CYAN
    p2 = tf_ai.add_paragraph()
    p2.text = "Biometric Verification & Landmark Matching"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.name = FONT_NAME
    p2.font.size = Pt(8.5)
    p2.font.color.rgb = TEXT_MUTED
    
    # Arrow from CompreFace down to Result
    arr_res = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(9.4), Inches(4.20), Inches(0.28), Inches(0.14))
    arr_res.fill.solid()
    arr_res.fill.fore_color.rgb = ACCENT_GREEN
    arr_res.line.fill.background()
    
    # Attendance Result / Report Box
    b_res = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.333), Inches(4.36), Inches(4.5), Inches(0.56))
    b_res.fill.solid()
    b_res.fill.fore_color.rgb = CARD_INNER
    b_res.line.color.rgb = ACCENT_GREEN
    b_res.line.width = Pt(1.5)
    tf_res = b_res.text_frame
    p = tf_res.paragraphs[0]
    p.text = "Attendance Result / Live Report Output"
    p.alignment = PP_ALIGN.CENTER
    p.font.name = FONT_NAME
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN
    p2 = tf_res.add_paragraph()
    p2.text = "Instant Status Confirmation & Percentage Updates"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.name = FONT_NAME
    p2.font.size = Pt(8.5)
    p2.font.color.rgb = TEXT_MUTED

    # Bottom Area: Deployment Topology Card
    dep_card = add_card(slide, Inches(0.8), Inches(5.42), Inches(11.733), Inches(1.38))
    dep_card.line.color.rgb = ACCENT_CYAN
    tf_d = dep_card.text_frame
    tf_d.word_wrap = True
    tf_d.margin_left = Pt(16)
    tf_d.margin_top = Pt(8)
    
    p_dh = tf_d.paragraphs[0]
    p_dh.text = "DEPLOYMENT ARCHITECTURE & CLOUD INTEGRATION"
    p_dh.font.name = FONT_NAME
    p_dh.font.size = Pt(11)
    p_dh.font.bold = True
    p_dh.font.color.rgb = ACCENT_CYAN
    p_dh.space_after = Pt(3)
    
    p_dt = tf_d.add_paragraph()
    p_dt.text = "Vercel (Global Edge Frontend)  ➔  Render Cloud (PHP Backend API)  ➔  Managed PostgreSQL & CompreFace Service"
    p_dt.font.name = FONT_NAME
    p_dt.font.size = Pt(11)
    p_dt.font.bold = True
    p_dt.font.color.rgb = TEXT_LIGHT
    p_dt.space_after = Pt(2)
    
    p_sub = tf_d.add_paragraph()
    p_sub.text = "Two-tier decoupled architecture: Static HTML/JS served via global CDN edge for zero latency; PHP API securely handles auth, attendance rules, and database queries over HTTPS."
    p_sub.font.name = FONT_NAME
    p_sub.font.size = Pt(9.5)
    p_sub.font.color.rgb = TEXT_MUTED

# -----------------------------------------------------------------------------
# SLIDE 10: 9. APPLICATIONS
# -----------------------------------------------------------------------------
def build_slide_10():
    slide = create_base_slide()
    add_header(slide, "9. Applications", "9. Applications", "Target Institutional Deployments & Operational Scenarios")
    
    apps = [
        ("Schools & Junior Colleges", 
         "Streamlining daily morning attendance and tracking student attendance regularity.", ACCENT_CYAN),
        ("Polytechnic Institutes", 
         "Department-wise (CO, IF, EJ, ME) lecture and laboratory session roll-call automation.", ACCENT_BLUE),
        ("Engineering Colleges & Universities", 
         "Scalable tracking across multiple classes, divisions, and semester curricula.", ACCENT_INDIGO),
        ("Classroom Lecture Attendance", 
         "Rapid attendance marking in theory lectures saving 10–15 minutes per period.", ACCENT_GREEN),
        ("Practical / Laboratory Batches", 
         "Sub-division batch-wise (B1, B2, B3) tracking for technical experiments.", ACCENT_CYAN),
        ("Department-Wise Management", 
         "Consolidated student directories, faculty allocations, and class timetable oversight.", ACCENT_BLUE),
        ("Student Attendance Reporting", 
         "Transparent percentage tracking, subject breakdown, and monthly attendance calendars.", ACCENT_INDIGO),
        ("Teacher Attendance Workflows", 
         "Structured session creation, editable records, and audited override mechanisms.", ACCENT_GREEN),
        ("Centralized Academic Records", 
         "Secure digital persistence replacing vulnerable paper attendance registers.", ACCENT_CYAN)
    ]
    
    cols = 3
    col_w = Inches(3.75)
    row_h = Inches(1.55)
    start_x = Inches(0.8)
    start_y = Inches(1.6)
    gap_x = Inches(0.24)
    gap_y = Inches(0.2)
    
    for idx, (title, desc, color) in enumerate(apps):
        r = idx // cols
        c = idx % cols
        x = start_x + c * (col_w + gap_x)
        y = start_y + r * (row_h + gap_y)
        
        card = add_card(slide, x, y, col_w, row_h)
        card.line.color.rgb = color
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Pt(12)
        tf.margin_top = Pt(10)
        tf.margin_right = Pt(12)
        
        p = tf.paragraphs[0]
        p.text = "✦  " + title
        p.font.name = FONT_NAME
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = color
        p.space_after = Pt(3)
        
        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = FONT_NAME
        p_desc.font.size = Pt(9.5)
        p_desc.font.color.rgb = TEXT_LIGHT

# -----------------------------------------------------------------------------
# SLIDE 11: 10. ADVANTAGES AND LIMITATIONS
# -----------------------------------------------------------------------------
def build_slide_11():
    slide = create_base_slide()
    add_header(slide, "10. Advantages & Limitations", "10. Advantages & Limitations", "System Evaluation, Technical Benefits, and Real-World Constraints")
    
    card_adv = add_card(slide, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.1))
    card_adv.line.color.rgb = ACCENT_GREEN
    tf_adv = card_adv.text_frame
    tf_adv.word_wrap = True
    tf_adv.margin_left = Pt(16)
    tf_adv.margin_top = Pt(14)
    tf_adv.margin_right = Pt(16)
    
    p_h = tf_adv.paragraphs[0]
    p_h.text = "ADVANTAGES"
    p_h.font.name = FONT_NAME
    p_h.font.size = Pt(13)
    p_h.font.bold = True
    p_h.font.color.rgb = ACCENT_GREEN
    p_h.space_after = Pt(8)
    
    advantages = [
        "Significantly reduces teacher manual effort during daily classroom sessions.",
        "Saves valuable instructional lecture time (recovering 10–15 minutes).",
        "Centralized digital records eliminate paper register loss and damage.",
        "Rapid attendance processing through batch marking or AI scanning.",
        "Eliminates human arithmetic calculation and transcription errors.",
        "Provides comprehensive student, faculty, and academic directory management.",
        "Automated attendance reports, analytics, and 75% shortage alert rosters.",
        "Supports AI-based face verification for modern contactless attendance.",
        "Strict role-based access control (Admin, Teacher, Student) and audit logs.",
        "Universal accessibility through any modern web browser across campus."
    ]
    for adv in advantages:
        p = tf_adv.add_paragraph()
        p.text = "✔  " + adv
        p.font.name = FONT_NAME
        p.font.size = Pt(9.8)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(2.5)

    card_lim = add_card(slide, Inches(6.8), Inches(1.6), Inches(5.733), Inches(5.1))
    card_lim.line.color.rgb = ACCENT_AMBER
    tf_lim = card_lim.text_frame
    tf_lim.word_wrap = True
    tf_lim.margin_left = Pt(16)
    tf_lim.margin_top = Pt(14)
    tf_lim.margin_right = Pt(16)
    
    p_lh = tf_lim.paragraphs[0]
    p_lh.text = "LIMITATIONS & CONSTRAINTS"
    p_lh.font.name = FONT_NAME
    p_lh.font.size = Pt(13)
    p_lh.font.bold = True
    p_lh.font.color.rgb = ACCENT_AMBER
    p_lh.space_after = Pt(8)
    
    limitations = [
        "Requires a functional webcam or built-in camera for face attendance.",
        "Face recognition accuracy depends on ambient lighting & image clarity.",
        "Cloud deployment requires active campus internet/network connectivity.",
        "AI verification requires prior enrollment of registered student face data.",
        "Initial cloud deployment and Docker setup require technical administration.",
        "Biometric data handling requires appropriate data security and user privacy controls.",
        "Face recognition may occasionally fail; manual teacher fallback is mandatory."
    ]
    for lim in limitations:
        p = tf_lim.add_paragraph()
        p.text = "✖  " + lim
        p.font.name = FONT_NAME
        p.font.size = Pt(10)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(5)
        
    p_note = tf_lim.add_paragraph()
    p_note.text = "Note: SAMS guarantees full academic integrity by keeping manual roll-call always available as a primary fallback."
    p_note.font.name = FONT_NAME
    p_note.font.size = Pt(9)
    p_note.font.italic = True
    p_note.font.color.rgb = TEXT_MUTED
    p_note.space_before = Pt(8)

# -----------------------------------------------------------------------------
# SLIDE 12: 11. FUTURE SCOPE
# -----------------------------------------------------------------------------
def build_slide_12():
    slide = create_base_slide()
    add_header(slide, "11. Future Scope", "11. Future Scope", "Planned Enhancements & Advanced Research Directions")
    
    future_items = [
        ("Improved Face Recognition Accuracy", 
         "Fine-tuning deep learning models with edge adaptive lighting correction for higher multi-face precision.", ACCENT_CYAN),
        ("Multi-Frame Identity Confirmation", 
         "Capturing short video burst sequences to vote across multiple consecutive frames for high confidence.", ACCENT_BLUE),
        ("Advanced Liveness & Anti-Spoofing", 
         "Implementing dynamic blink detection, head motion vectors, and planar glare analysis to prevent photo spoofing.", ACCENT_INDIGO),
        ("Native Mobile Application", 
         "Developing cross-platform Flutter/React Native mobile applications for on-the-go teacher attendance.", ACCENT_GREEN),
        ("Advanced Attendance Analytics", 
         "Machine learning predictive models to forecast dropout risks and attendance deficit trends early.", ACCENT_CYAN),
        ("AI-Generated Insights", 
         "Automated natural language summaries for institute principals and HODs highlighting attendance KPIs.", ACCENT_BLUE),
        ("Parent & Student Notifications", 
         "Automated WhatsApp, SMS, and Email alert triggers whenever student attendance drops below 75%.", ACCENT_INDIGO),
        ("QR-Based Attendance Option", 
         "Generating dynamic time-synchronized encrypted QR codes for supplementary high-speed verification.", ACCENT_GREEN),
        ("College ERP Integration", 
         "Direct API integration with State Board (MSBTE) examination portals and college ERP systems.", ACCENT_CYAN)
    ]
    
    cols = 3
    col_w = Inches(3.75)
    row_h = Inches(1.55)
    start_x = Inches(0.8)
    start_y = Inches(1.6)
    gap_x = Inches(0.24)
    gap_y = Inches(0.2)
    
    for idx, (title, desc, color) in enumerate(future_items):
        r = idx // cols
        c = idx % cols
        x = start_x + c * (col_w + gap_x)
        y = start_y + r * (row_h + gap_y)
        
        card = add_card(slide, x, y, col_w, row_h)
        card.line.color.rgb = color
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Pt(12)
        tf.margin_top = Pt(10)
        tf.margin_right = Pt(12)
        
        p = tf.paragraphs[0]
        p.text = "FUTURE: " + title
        p.font.name = FONT_NAME
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = color
        p.space_after = Pt(3)
        
        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.name = FONT_NAME
        p_desc.font.size = Pt(9)
        p_desc.font.color.rgb = TEXT_LIGHT

# -----------------------------------------------------------------------------
# SLIDE 13: 12. EXPECTED OUTCOME
# -----------------------------------------------------------------------------
def build_slide_13():
    slide = create_base_slide()
    add_header(slide, "12. Expected Outcome", "12. Expected Outcome", "Project Deliverables & Institutional Value Addition")
    
    trans_banner = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.55), Inches(11.733), Inches(1.0))
    trans_banner.fill.solid()
    trans_banner.fill.fore_color.rgb = CARD_INNER
    trans_banner.line.color.rgb = ACCENT_CYAN
    trans_banner.line.width = Pt(1.5)
    
    tf_tb = trans_banner.text_frame
    tf_tb.word_wrap = True
    tf_tb.margin_top = Pt(12)
    p_t = tf_tb.paragraphs[0]
    p_t.alignment = PP_ALIGN.CENTER
    
    r1 = p_t.add_run()
    r1.text = "Manual Paper Registers"
    r1.font.name = FONT_NAME
    r1.font.size = Pt(14)
    r1.font.bold = True
    r1.font.color.rgb = ACCENT_AMBER
    
    r_arr1 = p_t.add_run()
    r_arr1.text = "   ➔   "
    r_arr1.font.name = FONT_NAME
    r_arr1.font.size = Pt(14)
    r_arr1.font.bold = True
    r_arr1.font.color.rgb = ACCENT_CYAN
    
    r2 = p_t.add_run()
    r2.text = "SAMS Centralized Platform"
    r2.font.name = FONT_NAME
    r2.font.size = Pt(14)
    r2.font.bold = True
    r2.font.color.rgb = ACCENT_BLUE
    
    r_arr2 = p_t.add_run()
    r_arr2.text = "   ➔   "
    r_arr2.font.name = FONT_NAME
    r_arr2.font.size = Pt(14)
    r_arr2.font.bold = True
    r_arr2.font.color.rgb = ACCENT_CYAN
    
    r3 = p_t.add_run()
    r3.text = "Automated Digital Attendance"
    r3.font.name = FONT_NAME
    r3.font.size = Pt(14)
    r3.font.bold = True
    r3.font.color.rgb = ACCENT_GREEN

    outcomes = [
        ("Digitalization & Process Efficiency", 
         ["Completely digitalizes institutional attendance tracking across departments.",
          "Eliminates repetitive physical paperwork and manual roll-call overhead.",
          "Saves 10–15 minutes per lecture, returning critical instructional time to faculty."],
         ACCENT_CYAN, Inches(0.8), Inches(2.75), Inches(5.7), Inches(2.0)),
         
        ("Data Centralization & Accuracy", 
         ["Centralizes all academic records in a high-performance PostgreSQL database.",
          "Eliminates transcription errors, proxy attendance, and lost register risks.",
          "Enforces unique constraints preventing duplicate session markings."],
         ACCENT_INDIGO, Inches(6.8), Inches(2.75), Inches(5.733), Inches(2.0)),
         
        ("Practical Emerging AI Application", 
         ["Demonstrates practical deployment of AI and Computer Vision in education.",
          "Explores biometric vector matching and server-side face verification.",
          "Integrates modern camera streaming via standard web browser APIs."],
         ACCENT_GREEN, Inches(0.8), Inches(4.95), Inches(5.7), Inches(1.9)),
         
        ("Actionable Insights & Compliance", 
         ["Provides instant visibility into individual student and subject attendance.",
          "Automated shortage detection flags students below the mandatory 75% threshold.",
          "Establishes a solid, scalable foundation for future intelligent campus features."],
         ACCENT_BLUE, Inches(6.8), Inches(4.95), Inches(5.733), Inches(1.9))
    ]
    
    for title, points, color, left, top, width, height in outcomes:
        card = add_card(slide, left, top, width, height)
        card.line.color.rgb = color
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Pt(14)
        tf.margin_top = Pt(10)
        tf.margin_right = Pt(14)
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = FONT_NAME
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = color
        p.space_after = Pt(4)
        
        for pt in points:
            p_pt = tf.add_paragraph()
            p_pt.text = "• " + pt
            p_pt.font.name = FONT_NAME
            p_pt.font.size = Pt(9.5)
            p_pt.font.color.rgb = TEXT_LIGHT
            p_pt.space_after = Pt(2)

# -----------------------------------------------------------------------------
# SLIDE 14: 13. CONCLUSION
# -----------------------------------------------------------------------------
def build_slide_14():
    slide = create_base_slide()
    add_header(slide, "13. Conclusion", "13. Conclusion", "Summary & Acknowledgments")
    
    # Left Card: Academic Summary
    card_l = add_card(slide, Inches(0.8), Inches(1.6), Inches(7.2), Inches(5.1))
    card_l.line.color.rgb = ACCENT_CYAN
    tf_l = card_l.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = Pt(18)
    tf_l.margin_top = Pt(16)
    tf_l.margin_right = Pt(18)
    
    p_h = tf_l.paragraphs[0]
    p_h.text = "PROJECT SUMMARY & TAKEAWAYS"
    p_h.font.name = FONT_NAME
    p_h.font.size = Pt(13)
    p_h.font.bold = True
    p_h.font.color.rgb = ACCENT_CYAN
    p_h.space_after = Pt(10)
    
    conclusions = [
        ("Centralized Solution", "SAMS provides a comprehensive, centralized digital solution for modern student attendance management."),
        ("Emerging Tech Synergy", "Successfully combines responsive web technologies, PostgreSQL database management, and emerging AI computer vision."),
        ("Automated Verification", "Face recognition automates student identity verification while reducing dependence on tedious manual roll-calls."),
        ("Operational Reliability", "Maintains high operational reliability with instantaneous manual roll-call fallback for uninterrupted classroom flow."),
        ("Practical Problem Solving", "Demonstrates how modern emerging technologies can be applied effectively to solve practical educational administration challenges."),
        ("Extensible Foundation", "Establishes a solid, scalable architectural foundation for further development in intelligent academic campus management systems.")
    ]
    
    for title, desc in conclusions:
        p = tf_l.add_paragraph()
        p.text = f"• {title}: "
        p.font.name = FONT_NAME
        p.font.size = Pt(10.5)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE
        
        run = p.add_run()
        run.text = desc
        run.font.name = FONT_NAME
        run.font.size = Pt(10)
        run.font.bold = False
        run.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(6)

    # Right Card: Professional Thank You & Presentation Details
    card_r = add_card(slide, Inches(8.3), Inches(1.6), Inches(4.233), Inches(5.1))
    card_r.line.color.rgb = ACCENT_GREEN
    
    # Large Thank You Box inside card_r
    ty_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.6), Inches(1.85), Inches(3.633), Inches(1.4))
    ty_box.fill.solid()
    ty_box.fill.fore_color.rgb = CARD_INNER
    ty_box.line.color.rgb = ACCENT_GREEN
    ty_box.line.width = Pt(1.5)
    
    tf_ty = ty_box.text_frame
    tf_ty.word_wrap = True
    tf_ty.margin_top = Pt(12)
    p_ty = tf_ty.paragraphs[0]
    p_ty.text = "THANK YOU!"
    p_ty.alignment = PP_ALIGN.CENTER
    p_ty.font.name = FONT_NAME
    p_ty.font.size = Pt(28)
    p_ty.font.bold = True
    p_ty.font.color.rgb = ACCENT_GREEN
    
    p_subty = tf_ty.add_paragraph()
    p_subty.text = "Questions & Discussion"
    p_subty.alignment = PP_ALIGN.CENTER
    p_subty.font.name = FONT_NAME
    p_subty.font.size = Pt(11)
    p_subty.font.bold = True
    p_subty.font.color.rgb = ACCENT_CYAN
    p_subty.space_before = Pt(4)
    
    # Metadata Box below Thank You
    det_box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.6), Inches(3.45), Inches(3.633), Inches(3.05))
    det_box.fill.solid()
    det_box.fill.fore_color.rgb = CARD_INNER
    det_box.line.color.rgb = CARD_BORDER
    
    lbl_b = slide.shapes.add_textbox(Inches(8.8), Inches(3.6), Inches(1.5), Inches(2.7))
    tf_lb = lbl_b.text_frame
    tf_lb.word_wrap = True
    tf_lb.margin_left = 0
    tf_lb.margin_top = 0
    
    val_b = slide.shapes.add_textbox(Inches(10.3), Inches(3.6), Inches(1.8), Inches(2.7))
    tf_vb = val_b.text_frame
    tf_vb.word_wrap = True
    tf_vb.margin_left = 0
    tf_vb.margin_top = 0
    
    meta_info = [
        ("Presented by:", "Vishal Yadav"),
        ("Class / Branch:", "CO5K (Computer)"),
        ("Enrollment No.:", "________________"),
        ("Subject:", "SPI (Emerging Trends)"),
        ("Project:", "SAMS Platform"),
        ("Status:", "SPI Completed")
    ]
    for i, (k, v) in enumerate(meta_info):
        p_k = tf_lb.paragraphs[0] if i == 0 else tf_lb.add_paragraph()
        p_k.text = k
        p_k.font.name = FONT_NAME
        p_k.font.size = Pt(10)
        p_k.font.bold = True
        p_k.font.color.rgb = ACCENT_CYAN
        p_k.space_after = Pt(4)
        
        p_v = tf_vb.paragraphs[0] if i == 0 else tf_vb.add_paragraph()
        p_v.text = v
        p_v.font.name = FONT_NAME
        p_v.font.size = Pt(10)
        p_v.font.bold = (k == "Presented by:" or k == "Enrollment No.:")
        p_v.font.color.rgb = TEXT_LIGHT
        p_v.space_after = Pt(4)

# -----------------------------------------------------------------------------
# SLIDE 15: 14. REFERENCES
# -----------------------------------------------------------------------------
def build_slide_15():
    slide = create_base_slide()
    add_header(slide, "14. References", "14. References", "Official Technical Documentation & Academic Resources")
    
    card_l = add_card(slide, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.1))
    card_l.line.color.rgb = ACCENT_CYAN
    tf_l = card_l.text_frame
    tf_l.word_wrap = True
    tf_l.margin_left = Pt(16)
    tf_l.margin_top = Pt(14)
    tf_l.margin_right = Pt(16)
    
    p_lh = tf_l.paragraphs[0]
    p_lh.text = "OFFICIAL TECHNOLOGY DOCUMENTATION"
    p_lh.font.name = FONT_NAME
    p_lh.font.size = Pt(12)
    p_lh.font.bold = True
    p_lh.font.color.rgb = ACCENT_CYAN
    p_lh.space_after = Pt(8)
    
    refs_tech = [
        ("Exadel CompreFace Documentation", "Open-source facial recognition service via REST API.\nhttps://github.com/exadel-inc/CompreFace"),
        ("PHP 8.2 Official Manual", "Server-side scripting language & PDO database abstraction.\nhttps://www.php.net/docs.php"),
        ("PostgreSQL 13+ Relational Database", "Official relational database system manual & documentation.\nhttps://www.postgresql.org/docs/"),
        ("MDN Web Docs (Mozilla Developer Network)", "JavaScript standards, Fetch API, and MediaDevices camera stream.\nhttps://developer.mozilla.org/"),
        ("Vercel Hosting Documentation", "Global edge static hosting & routing configuration.\nhttps://vercel.com/docs"),
        ("Render Cloud Application Platform", "Docker container deployment & managed PostgreSQL services.\nhttps://render.com/docs")
    ]
    for title, desc in refs_tech:
        p = tf_l.add_paragraph()
        p.text = f"• {title}\n  {desc}"
        p.font.name = FONT_NAME
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(5)

    card_r = add_card(slide, Inches(6.8), Inches(1.6), Inches(5.733), Inches(5.1))
    card_r.line.color.rgb = ACCENT_BLUE
    tf_r = card_r.text_frame
    tf_r.word_wrap = True
    tf_r.margin_left = Pt(16)
    tf_r.margin_top = Pt(14)
    tf_r.margin_right = Pt(16)
    
    p_rh = tf_r.paragraphs[0]
    p_rh.text = "ACADEMIC & INDUSTRY STANDARDS"
    p_rh.font.name = FONT_NAME
    p_rh.font.size = Pt(12)
    p_rh.font.bold = True
    p_rh.font.color.rgb = ACCENT_BLUE
    p_rh.space_after = Pt(8)
    
    refs_academic = [
        ("MSBTE Curriculum Manual", "Maharashtra State Board of Technical Education guidelines for Seminar & Project Initiation (SPI) in Computer Engineering (CO5K)."),
        ("GitHub Version Control Documentation", "Collaborative development, git branching, and repository management.\nhttps://docs.github.com/"),
        ("Computer Vision & Deep Learning Literature", "Research publications on Deep Facial Recognition, landmark vector representations, and convolutional neural networks."),
        ("OWASP Web Security Standards", "Open Web Application Security Project guidelines for session management, prepared SQL queries, and input validation."),
        ("Cloud Computing in Education", "Academic literature and IEEE conference papers on digital campus automation and attendance management systems.")
    ]
    for title, desc in refs_academic:
        p = tf_r.add_paragraph()
        p.text = f"• {title}\n  {desc}"
        p.font.name = FONT_NAME
        p.font.size = Pt(9.5)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(6)

# -----------------------------------------------------------------------------
# BUILD ALL SLIDES AND SAVE
# -----------------------------------------------------------------------------
def main():
    print("Rebuilding SAMS SPI Project Presentation...")
    build_slide_1()
    build_slide_2()
    build_slide_3()
    build_slide_4()
    build_slide_5()
    build_slide_6()
    build_slide_7()
    build_slide_8()
    build_slide_9()
    build_slide_10()
    build_slide_11()
    build_slide_12()
    build_slide_13()
    build_slide_14()
    build_slide_15()
    
    output_filename = "SAMS_SPI_Project_Presentation_Vishal_Yadav.pptx"
    prs.save(output_filename)
    print(f"Presentation successfully created locally: {output_filename}")

    # Synchronize presentation to Desktop
    desktop_dirs = [
        os.path.join(os.environ.get("USERPROFILE", ""), "OneDrive", "Desktop"),
        os.path.join(os.environ.get("USERPROFILE", ""), "Desktop")
    ]
    saved_to_desktop = False
    for d in desktop_dirs:
        if os.path.exists(d):
            dest = os.path.join(d, output_filename)
            prs.save(dest)
            print(f"Presentation successfully saved to Desktop: {dest}")
            saved_to_desktop = True

    if not saved_to_desktop:
        print("Note: Could not find Desktop folder automatically.")

    print(f"Total slides generated: {len(prs.slides)}")

if __name__ == "__main__":
    main()
